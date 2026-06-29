"""
analyze_spending.py
Reads the SBE travel workbook, applies the Best SBE Cost fallback logic
used in the Master tab, and prints summary analysis matching the Excel pivots.
"""

import pandas as pd
from openpyxl import load_workbook

SOURCE_FILE = "SBE Travel Spreadsheet working.xlsx"

# Rows to exclude from analysis (these are budget allocations, not real trips)
EXCLUDED_NAMES = ["ALLOCATED", "BALANCES", "Totals"]

# Budgets pulled from the title row of each yearly tab
BUDGETS = {
    "2018-2019": 23400,
    "2019-2020": 25000,
    "2020-2021": None,        # COVID year, no budget set
    "2021-2022": 12383,
    "2022-2023": 12383,
    "2023-2024": 12383,
    "2024-2025": 23394,
    "2025-2026": 24355,
}


def load_master():
    """Load the Master tab into a DataFrame."""
    wb = load_workbook(SOURCE_FILE, data_only=True)
    ws = wb["Master"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    data = []
    for r in range(2, ws.max_row + 1):
        row = {h: ws.cell(row=r, column=c).value for c, h in enumerate(headers, start=1)}
        data.append(row)
    return pd.DataFrame(data)


def calculate_best_cost(row):
    """Mirror of the Excel Best SBE Cost formula:
       Total SBE Expense > Prepaid + Paid to Traveler > Approved > Total Expenses"""
    if pd.notna(row["Total SBE Expense"]):
        return row["Total SBE Expense"]
    prepaid = row["Prepaid Expenses"] if pd.notna(row["Prepaid Expenses"]) else 0
    paid = row["Paid to Traveler"] if pd.notna(row["Paid to Traveler"]) else 0
    if prepaid + paid > 0:
        return prepaid + paid
    if pd.notna(row["Approved Amount"]):
        return row["Approved Amount"]
    if pd.notna(row["Total Expenses"]):
        return row["Total Expenses"]
    return None


def filter_real_trips(df):
    """Remove allocation rows, balances, totals, and rows flagged as duplicates."""
    df = df[~df["Traveler Name"].isin(EXCLUDED_NAMES)]
    df = df[~df["Notes"].astype(str).str.startswith("DUPLICATE", na=False)]
    return df


def main():
    print(f"Loading {SOURCE_FILE}...")
    df = load_master()

    money_cols = ["Total Expenses", "Approved Amount", "Prepaid Expenses",
                  "Paid to Traveler", "Other Funding", "Total SBE Expense"]
    for col in money_cols:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    df["Best SBE Cost"] = df.apply(calculate_best_cost, axis=1)
    df = filter_real_trips(df)

    print(f"\nAnalyzing {len(df)} real trips (excluded allocations/totals/duplicates)")
    print(f"Total SBE spend (8 years): ${df['Best SBE Cost'].sum():,.2f}")
    print(f"Average trip cost: ${df['Best SBE Cost'].mean():,.2f}")
    print(f"Median trip cost: ${df['Best SBE Cost'].median():,.2f}")

    print("\n=== Spending by Fiscal Year ===")
    by_year = df.groupby("Fiscal Year").agg(
        trips=("Best SBE Cost", "count"),
        total_spent=("Best SBE Cost", "sum")
    ).round(2)
    print(by_year.to_string())

    print("\n=== Budget vs. Actual ===")
    for fy, budget in BUDGETS.items():
        actual = df[df["Fiscal Year"] == fy]["Best SBE Cost"].sum()
        if budget:
            variance = actual - budget
            variance_pct = (variance / budget) * 100
            status = "OVER" if variance > 0 else "UNDER"
            print(f"  {fy}: Budget ${budget:>8,.0f} | Actual ${actual:>10,.2f} | "
                  f"Variance ${variance:>+10,.2f} ({variance_pct:+.1f}% {status})")
        else:
            print(f"  {fy}: Budget  NOT SET  | Actual ${actual:>10,.2f}")

    print("\n=== Spending by Travel Purpose ===")
    by_purpose = df.groupby("Travel Purpose").agg(
        trips=("Best SBE Cost", "count"),
        total_spent=("Best SBE Cost", "sum")
    ).round(2).sort_values("total_spent", ascending=False)
    print(by_purpose.to_string())

    print("\n=== Top 10 Travelers by Spending ===")
    by_traveler = df.groupby("Traveler Name").agg(
        trips=("Best SBE Cost", "count"),
        total_spent=("Best SBE Cost", "sum")
    ).round(2).sort_values("total_spent", ascending=False).head(10)
    print(by_traveler.to_string())


if __name__ == "__main__":
    main()
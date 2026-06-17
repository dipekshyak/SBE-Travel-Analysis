"""
combine_tabs.py
Reads the SBE travel voucher workbook, combines all 8 fiscal year tabs into
one cleaned DataFrame, and saves the result to a new Excel file.
"""

import pandas as pd
from openpyxl import load_workbook
import re

# Configuration
SOURCE_FILE = "SBE Travel Spreadsheet working.xlsx"
OUTPUT_FILE = "combined_travel_data.xlsx"
YEAR_TABS = ["2018-19", "2019-20", "2020-21", "2021-22",
             "2022-23", "2023-24", "2024-25", "2025-26"]
OLD_FORMAT_TABS = ["2018-19", "2019-20", "2020-21", "2021-22"]


def fiscal_year_label(tab):
    """Convert tab name '2018-19' into long form '2018-2019'."""
    parts = tab.split("-")
    return f"{parts[0]}-20{parts[1]}"


def parse_dates(date_val):
    """Convert various date formats into (start_date, end_date) MM/DD/YYYY strings."""
    if date_val is None or str(date_val).strip() == "":
        return ("", "")
    if hasattr(date_val, 'strftime'):
        d = date_val.strftime("%m/%d/%Y")
        return (d, d)
    s = str(date_val).strip()
    # Date range like "8/10-8/14/18"
    range_match = re.match(r'(\d+)/(\d+)-(\d+)/(\d+)/(\d+)', s)
    if range_match:
        m1, d1, m2, d2, y = range_match.groups()
        yy = int(y)
        year = 2000 + yy if yy < 50 else 1900 + yy
        return (f"{int(m1):02d}/{int(d1):02d}/{year}",
                f"{int(m2):02d}/{int(d2):02d}/{year}")
    # Single date like "7/12/2018"
    single = re.match(r'(\d+)/(\d+)/(\d+)', s)
    if single:
        m, d, y = single.groups()
        yy = int(y)
        year = 2000 + yy if yy < 50 else 1900 + yy if yy < 100 else yy
        return (f"{int(m):02d}/{int(d):02d}/{year}",
                f"{int(m):02d}/{int(d):02d}/{year}")
    return (s, s)


def clean_money(val):
    """Strip $, commas, and convert to float. Return None if not a number."""
    if val is None or str(val).strip() == "":
        return None
    s = str(val).strip().replace("$", "").replace(",", "")
    num_match = re.match(r'(-?\d+\.?\d*)', s)
    if num_match:
        try:
            return float(num_match.group(1))
        except ValueError:
            return None
    return None


def clean_account(val):
    """Clean account numbers (remove trailing .0 from floats)."""
    if val is None:
        return ""
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s


def clean_name(name):
    """Standardize name to 'Last, First' with proper capitalization."""
    if not name or not str(name).strip():
        return ""
    s = str(name).strip()
    # Strip PJ# patterns
    s = re.sub(r'\(?\s*PJ\s*#?\s*\d+\s*\)?', '', s, flags=re.IGNORECASE).strip()
    # Fix missing space after comma
    s = re.sub(r',(\S)', r', \1', s)
    # Already Last, First format
    if ',' in s:
        parts = [p.strip() for p in s.split(',', 1)]
        if len(parts) == 2 and parts[0] and parts[1]:
            return f"{parts[0].title()}, {parts[1].title()}"
        return s.title()
    # First Last format — convert to Last, First
    parts = s.split()
    if len(parts) >= 2:
        last = parts[-1].title()
        first = " ".join(p.title() for p in parts[:-1])
        return f"{last}, {first}"
    return s.title()


def categorize_purpose(reason, name):
    """Best-guess Travel Purpose based on Reason for Travel."""
    if not reason or not str(reason).strip():
        return "Unspecified"
    r = str(reason).lower().strip()
    n = str(name).lower() if name else ""

    if "cancelled" in r or "repayment" in r or "non employee reimbursement" in r:
        return "Other / Administrative"
    if "candidate" in n or "interview" in r or "recruitment" in r:
        return "Recruitment"
    if "accepted student" in r:
        return "Recruitment"
    student_kw = ["student", "educational trip", "accompany", "ed trip", "w/students",
                  "w/ students", "harvard undergrad", "ibm site visit", "kpmg and scs"]
    if any(kw in r for kw in student_kw):
        return "Student Travel"
    if "branch campus" in r or ("teach" in r and "course" in r):
        return "Branch Campus Teaching"
    if any(kw in r for kw in ["present", "presenting", "presentation @"]):
        return "Conference Presentation"
    conf_kw = ["conf", "conference", "symposium", "summit", "annual meeting",
               "annual mtg", "show", "informs", "user group", "society for",
               "academy of", "institute", "consortium", "forum", "seminar",
               "advisory board", "ceremony", "induction"]
    if any(kw in r for kw in conf_kw):
        return "Conference Attendance"
    dev_kw = ["workshop", "training", "professional development", "certification",
              "online summit", "sail institute", "chair academy", "retreat"]
    if any(kw in r for kw in dev_kw):
        return "Professional Development"
    if "meeting" in r or "report to suny" in r:
        return "Meetings"
    if "research" in r:
        return "Research Support"
    return "Other / Administrative"


def combine_all_tabs():
    """Main function: read all 8 yearly tabs, combine, clean, and return DataFrame."""
    wb = load_workbook(SOURCE_FILE, data_only=True)
    all_rows = []

    for tab in YEAR_TABS:
        ws = wb[tab]
        fy = fiscal_year_label(tab)
        is_old = tab in OLD_FORMAT_TABS

        # Start from row 3 (row 1 = budget title, row 2 = headers)
        for r in range(3, ws.max_row + 1):
            name = ws.cell(row=r, column=1).value
            if name is None or str(name).strip() == "":
                continue

            dest = ws.cell(row=r, column=2).value
            dates = ws.cell(row=r, column=3).value
            reason = ws.cell(row=r, column=4).value
            total_exp = ws.cell(row=r, column=5).value

            if is_old:
                approved = ws.cell(row=r, column=6).value
                prepaid = ws.cell(row=r, column=7).value
                paid_trav = ws.cell(row=r, column=8).value
                other_fund = ws.cell(row=r, column=9).value
                total_sbe = ws.cell(row=r, column=10).value
                account = ws.cell(row=r, column=11).value
            else:
                approved = ws.cell(row=r, column=7).value
                prepaid = ws.cell(row=r, column=8).value
                paid_trav = ws.cell(row=r, column=9).value
                other_fund = ws.cell(row=r, column=10).value
                total_sbe = ws.cell(row=r, column=11).value
                account = ws.cell(row=r, column=12).value

            start_date, end_date = parse_dates(dates)
            cleaned_name = clean_name(name)
            purpose = categorize_purpose(reason, name)
            account_clean = clean_account(account)

            all_rows.append({
                "Fiscal Year": fy,
                "Traveler Name": cleaned_name,
                "Destination": str(dest).strip() if dest else "",
                "Start Date": start_date,
                "End Date": end_date,
                "Reason for Travel": str(reason).strip() if reason else "",
                "Travel Purpose": purpose,
                "Expense Type": "Travel & Related",
                "Funding Source": account_clean,
                "Total Expenses": clean_money(total_exp),
                "Approved Amount": clean_money(approved),
                "Prepaid Expenses": clean_money(prepaid),
                "Paid to Traveler": clean_money(paid_trav),
                "Other Funding": clean_money(other_fund),
                "Total SBE Expense": clean_money(total_sbe),
                "Account Number": account_clean,
                "Notes": "",
            })

    return pd.DataFrame(all_rows)


def main():
    print(f"Reading {SOURCE_FILE}...")
    df = combine_all_tabs()

    print(f"\nCombined {len(df)} trips across {df['Fiscal Year'].nunique()} fiscal years.")
    print("\nTrips per fiscal year:")
    print(df['Fiscal Year'].value_counts().sort_index().to_string())

    print("\nTravel Purpose distribution:")
    print(df['Travel Purpose'].value_counts().to_string())

    df.to_excel(OUTPUT_FILE, index=False, sheet_name="Combined")
    print(f"\nSaved cleaned data to {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
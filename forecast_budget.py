"""
forecast_budget.py
Builds a three-scenario forecast for FY 2026-2027 SBE travel spending
based on historical actuals from the Master tab.
"""

import pandas as pd
from openpyxl import load_workbook

SOURCE_FILE = "SBE Travel Spreadsheet working.xlsx"
EXCLUDED_NAMES = ["ALLOCATED", "BALANCES", "Totals"]


def load_master():
    """Load the Master tab into a DataFrame."""
    wb = load_workbook(SOURCE_FILE, data_only=True)
    ws = wb["Master"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    data = []
    for r in range(2, ws.max_row + 1):
        row = {h: ws.cell(row=r, column=c).value for c, h in enumerate(headers, start=1)}
        data.append(row)
    return pd.DataFrame(data)


def calculate_best_cost(row):
    """Mirror of the Excel Best SBE Cost fallback logic."""
    if pd.notna(row["Total SBE Expense"]):
        return row["Total SBE Expense"]
    prepaid = row["Prepaid Expenses"] if pd.notna(row["Prepaid Expenses"]) else 0
    paid = row["Paid to Traveler"] if pd.notna(row["Paid to Traveler"]) else 0
    if prepaid + paid > 0:
        return prepaid + paid
    if pd.notna(row["Approved Amount"]):
        return row["Approved Amount"]
    if pd.notna(row["Total Expenses"]):
        return row["Total Expenses"]
    return None


def filter_real_trips(df):
    """Remove allocations, balances, totals, and flagged duplicate rows."""
    df = df[~df["Traveler Name"].isin(EXCLUDED_NAMES)]
    df = df[~df["Notes"].astype(str).str.startswith("DUPLICATE", na=False)]
    return df


def build_forecast(df):
    """Build 3 scenario forecasts for FY 26-27."""
    yearly = df.groupby("Fiscal Year")["Best SBE Cost"].sum().round(2)

    # Pull known actuals
    fy_2122 = yearly.get("2021-2022", 0)
    fy_2223 = yearly.get("2022-2023", 0)
    fy_2324 = yearly.get("2023-2024", 0)
    fy_2425 = yearly.get("2024-2025", 0)
    fy_2526 = yearly.get("2025-2026", 0)

    # Low case: constrained COVID-era average (FY 21-22, 22-23, 23-24)
    low_case = round((fy_2122 + fy_2223 + fy_2324) / 3, 2)

    # Base case: weighted average of most recent 3 years, recent-heavy
    base_case = round(fy_2324 * 0.2 + fy_2425 * 0.3 + fy_2526 * 0.5, 2)

    # High case: FY 25-26 pace continues
    high_case = round(fy_2526, 2)

    return {
        "LOW CASE (COVID-era avg)": low_case,
        "BASE CASE (weighted recent avg)": base_case,
        "HIGH CASE (FY 25-26 continues)": high_case,
        "FY 25-26 Actual (reference)": fy_2526,
    }


def main():
    print(f"Loading {SOURCE_FILE}...")
    df = load_master()

    money_cols = ["Total Expenses", "Approved Amount", "Prepaid Expenses",
                  "Paid to Traveler", "Other Funding", "Total SBE Expense"]
    for col in money_cols:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    df["Best SBE Cost"] = df.apply(calculate_best_cost, axis=1)
    df = filter_real_trips(df)

    print(f"\nAnalyzing {len(df)} real trips")
    print("\n=== Historical Actuals ===")
    yearly = df.groupby("Fiscal Year")["Best SBE Cost"].sum().round(2)
    print(yearly.to_string())

    print("\n=== FY 2026-2027 Forecast ===")
    forecast = build_forecast(df)
    for scenario, amount in forecast.items():
        vs_current = amount - forecast["FY 25-26 Actual (reference)"]
        pct = (vs_current / forecast["FY 25-26 Actual (reference)"] * 100
               if forecast["FY 25-26 Actual (reference)"] else 0)
        marker = "" if scenario == "FY 25-26 Actual (reference)" else f"  ({pct:+.1f}% vs FY 25-26)"
        print(f"  {scenario:<40} ${amount:>10,.2f}{marker}")

    print("\n=== Recommendation ===")
    base = forecast["BASE CASE (weighted recent avg)"]
    lower_target = base * 1.10
    upper_target = base * 1.15
    print(f"  Recommended budget range (base + 10-15% contingency):")
    print(f"  ${lower_target:>10,.2f}  to  ${upper_target:>10,.2f}")


if __name__ == "__main__":
    main()
    